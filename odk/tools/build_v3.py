"""Preserve the canonical ATI form; add MDS entity geography and ID evidence."""
from copy import copy
from pathlib import Path
import hashlib
import json

import openpyxl

repo = Path(__file__).resolve().parents[2]
source = repo / 'odk/ATI_Farmers_Profile_ODK_Form_v2.xlsx'
target = repo / 'odk/ATI_Farmers_Profile_ODK_Form_v3.xlsx'
book = openpyxl.load_workbook(source)
survey = book['survey']
headers = [cell.value for cell in survey[1]]
columns = {key: index + 1 for index, key in enumerate(headers) if key}


def insert(after, values):
    survey.insert_rows(after + 1)
    for col in range(1, survey.max_column + 1):
        survey.cell(after + 1, col)._style = copy(survey.cell(after, col)._style)
    for key, value in values.items():
        survey.cell(after + 1, columns[key], value)
    return after + 1


names = lambda: {survey.cell(row, columns['name']).value: row
                 for row in range(2, survey.max_row + 1)}
changes = []
for prefix in ('', 'hh_member_'):
    for level, parent in (('region', None), ('zone', 'region'), ('woreda', 'zone'), ('kebele', 'woreda')):
        name = prefix + level
        row = names()[name]
        survey.cell(row, columns['type'], f'select_one_from_file {level}.csv')
        survey.cell(row, columns['choice_filter'], f'parent=${{{prefix + parent}}}' if parent else None)
        survey.cell(row, columns['required'], 'yes' if level != 'kebele' else None)
        for property in ('code', 'name'):
            # Entity values are UUIDs. The registry receives code, not UUID/name.
            field = 'code' if property == 'code' else 'label'
            row = insert(row, {'type': 'calculate', 'name': name + '_' + property,
                'calculation': f"instance('{level}')/root/item[name=${{{name}}}]/{field}"})
        changes.append(name)

row = names()['national_rid_confirm']
row = insert(row, {'type': 'select_one oan_identity_document_type', 'name': 'identity_document_type',
    'label::English (en)': 'Type of identity document',
    'label::Amharic (am)': 'Type of identity document',
    'label::Afaan Oromoo (om)': 'Type of identity document', 'required': 'yes'})
insert(row, {'type': 'image', 'name': 'identity_document',
    'label::English (en)': 'Photograph the identity document',
    'label::Amharic (am)': 'Photograph the identity document',
    'label::Afaan Oromoo (om)': 'Photograph the identity document', 'required': 'yes',
    'parameters': 'max-pixels=1600',
    'hint::English (en)': 'This offline intake is unverified. An approver checks the ID evidence.'})
row = names()['crop_date']
insert(row, {'type': 'select_one oan_crop_season', 'name': 'crop_season',
    'label::English (en)': 'Crop season', 'label::Amharic (am)': 'Crop season',
    'label::Afaan Oromoo (om)': 'Crop season', 'required': 'yes'})
for name in ('national_uid', 'hh_member_national_uid'):
    row = names()[name]
    survey.cell(row, columns['constraint'], "regex(., '^([0-9]{12}|[0-9]{16})$')")
    survey.cell(row, columns['constraint_message::English (en)'], 'Enter the 12-digit FIN or 16-digit FAN, without spaces.')

# ODK's supported repeat-derived select stores the actual Parcel ID, not
# the repeat index. Existing business groups/field names stay in place.
for after_name, field, label in (
    ('crop_repeat', 'crop_land_id', 'Select the plot for this crop'),
    ('livestock_repeat', 'livestock_land_id', 'Select the plot for this livestock'),
):
    insert(names()[after_name], {'type': 'select_one ${land_id}', 'name': field,
        'label::English (en)': label, 'label::Amharic (am)': label,
        'label::Afaan Oromoo (om)': label, 'required': 'yes'})
# Input, machinery and finance answers are root sibling groups in the source.
# One explicit selection scopes their one native FarmInputs row.
insert(names()['agricultural_input'] - 1, {'type': 'select_one ${land_id}',
    'name': 'farm_input_land_id', 'required': 'yes',
    'relevant': "count(${land_id}[. != '']) > 0",
    'label::English (en)': 'Select the plot for the input, water, machinery and finance answers',
    'label::Amharic (am)': 'Select the plot for the input, water, machinery and finance answers',
    'label::Afaan Oromoo (om)': 'Select the plot for the input, water, machinery and finance answers'})

choices = book['choices']
choice_columns = {cell.value: i + 1 for i, cell in enumerate(choices[1]) if cell.value}
for list_name, options in (
    ('oan_identity_document_type', ('FAYDA', 'KEBELE_WOREDA_ID', 'PASSPORT', 'DRIVING_LICENCE',
        'KEBELE_CONFIRMATION_LETTER', 'OTHER_GOVERNMENT_ID')),
    ('oan_crop_season', ('SEASON_MEHER', 'SEASON_BELG', 'SEASON_IRRIGATION')),
):
    for option in options:
        row = choices.max_row + 1
        choices.cell(row, choice_columns['list_name'], list_name)
        choices.cell(row, choice_columns['name'], option)
        for key, col in choice_columns.items():
            if key.startswith('label'):
                choices.cell(row, col, option.replace('_', ' ').title())

settings = book['settings']
settings_columns = {cell.value: cell.column for cell in settings[1] if cell.value}
settings.cell(2, settings_columns['version'], 'v3')
book.save(target)
verify = openpyxl.load_workbook(target, read_only=True)
assert verify['settings'].cell(2, settings_columns['version']).value == 'v3'
assert all(name in names() for name in ('identity_document_type', 'identity_document', 'region_code', 'hh_member_kebele_code', 'crop_season'))
print(json.dumps({'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'output_sha256': hashlib.sha256(target.read_bytes()).hexdigest(), 'version': 'v3',
    'entity_location_questions': changes, 'published': False, 'path': str(target)}, indent=2))
