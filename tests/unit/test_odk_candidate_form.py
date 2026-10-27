"""Compile the unpublished source form and verify explicit parcel-value links."""
from pathlib import Path
import xml.etree.ElementTree as ET

import openpyxl
from pyxform.xls2xform import xls2xform_convert

ROOT = Path(__file__).resolve().parents[2]
NS = {'x': 'http://www.w3.org/2002/xforms'}


def records(book, sheet):
    rows = book[sheet].iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


def test_candidate_retains_source_questions_and_choices():
    source = openpyxl.load_workbook(ROOT/'odk/ATI_Farmers_Profile_ODK_Form_v2.xlsx', read_only=True)
    candidate = openpyxl.load_workbook(ROOT/'odk/ATI_Farmers_Profile_ODK_Form_v3.xlsx', read_only=True)
    original = {r['name'] for r in records(source, 'survey') if r.get('name')}
    current = {r['name'] for r in records(candidate, 'survey') if r.get('name')}
    assert original <= current
    old_choices = {(r['list_name'], r['name']) for r in records(source, 'choices') if r.get('name')}
    new_choices = {(r['list_name'], r['name']) for r in records(candidate, 'choices') if r.get('name')}
    assert old_choices <= new_choices


def test_generated_selects_store_parcel_ids_and_require_explicit_answers(tmp_path):
    xml = tmp_path/'farmer.xml'
    xls2xform_convert(str(ROOT/'odk/ATI_Farmers_Profile_ODK_Form_v3.xlsx'), str(xml), validate=False)
    root = ET.parse(xml).getroot()
    expected = {
        '/data/crop_information/crop_repeat/crop_land_id',
        '/data/livestock_info/livestock_repeat/livestock_land_id',
        '/data/farm_input_land_id',
    }
    controls = {n.attrib.get('ref'): n for n in root.findall('.//x:select1', NS)}
    binds = {n.attrib['nodeset']: n for n in root.findall('.//x:bind', NS)}
    for path in expected:
        itemset = controls[path].find('x:itemset', NS)
        assert itemset is not None
        assert itemset.attrib['nodeset'] == "/data/land_info/land_info_repeat[./land_id != '']"
        assert itemset.find('x:value', NS).attrib['ref'] == 'land_id'
        assert itemset.find('x:label', NS).attrib['ref'] == 'land_id'
        assert binds[path].attrib['required'] == 'true()'
    assert 'count(' in binds['/data/farm_input_land_id'].attrib['relevant']
